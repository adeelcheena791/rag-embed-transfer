"""Export the author version of the paper to Word (.docx) and plain text from the same LaTeX source.

The PDF built from paper/latex/main_author.tex is the authoritative version. This script flattens the LaTeX
(inputs, number macros, figure macros), resolves every \\ref from the LaTeX .aux file so section, table and
figure numbers match the PDF, and converts with pandoc (citations via --citeproc and references.bib).
Word output is single-column (Times New Roman 11 pt, A4, 2.5 cm margins, page numbers bottom centre).

    cd paper/latex && latexmk -pdf main_author.tex && cd ../.. && python -m src.paper.make_docx
"""
from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

LATEX = Path("paper/latex")
OUT = Path("paper/export")


def read_inputs(text: str) -> str:
    def repl(m):
        name = m.group(1)
        p = LATEX / (name if name.endswith(".tex") else name + ".tex")
        return read_inputs(p.read_text()) if p.exists() else ""
    return re.sub(r"\\input\{([^}]*)\}", repl, text)


def labels() -> dict[str, str]:
    aux = (LATEX / "main_author.aux").read_text()
    return {m.group(1): m.group(2) for m in re.finditer(r"\\newlabel\{([^}]*)\}\{\{([^}]*)\}", aux)}


def balanced(s: str, i: int) -> tuple[str, int]:
    """s[i] == '{' -> (content, index after closing brace)"""
    depth, j = 0, i
    while True:
        if s[j] == "{":
            depth += 1
        elif s[j] == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1


def expand_macro(s: str, name: str, nargs: int, fn) -> str:
    out, i = [], 0
    pat = "\\" + name + "{"
    while True:
        k = s.find(pat, i)
        if k < 0:
            out.append(s[i:])
            return "".join(out)
        out.append(s[i:k])
        j, args = k + len(name) + 1, []
        for _ in range(nargs):
            a, j = balanced(s, j)
            args.append(a)
        out.append(fn(*args))
        i = j


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    nums = json.load(open(LATEX / "numbers_index.json"))
    lab = labels()
    src = (LATEX / "main.tex").read_text()
    # figure macros defined in main.tex (\figsix etc.) -> their \fig{...} bodies
    for m in re.finditer(r"\\newcommand\{\\(fig[a-z]+)\}\{(\\fig\{.*)\}\n", src):
        src = src.replace("\\" + m.group(1) + "\n", m.group(2) + "\n")
    body = src[src.index("\\begin{document}") + len("\\begin{document}"):src.index("\\end{document}")]
    body = read_inputs(body)
    for m in re.finditer(r"\\newcommand\{\\(fig[a-z]+)\}\{(\\fig\{.*)\}\n", src):
        body = body.replace("\\" + m.group(1) + "\n", m.group(2) + "\n")
    body = body.replace("\\ifdefined\\authorversion", "").replace("\\fi\n", "\n")
    body = re.sub(r"\\bibliography\{[^}]*\}", "", body)
    body = body.replace("\\maketitle", "").replace("\\appendix", "")

    body = expand_macro(body, "NUM", 1, lambda k: nums["NUM"][k])
    body = expand_macro(body, "CI", 1, lambda k: nums["CI"][k])
    # \\ensuremath is left to pandoc, which handles it both inside and outside math

    counters = {"figure": 0, "table": 0}

    def fig(env, file, cap, label):
        n = lab.get(label, "?")
        return (f"\n\\begin{{figure}}\n\\centering\n\\includegraphics[width=6in]{{{(Path('paper/figures') / file).resolve()}.png}}\n"
                f"\\caption{{Figure {n}: {cap}}}\n\\end{{figure}}\n")
    body = expand_macro(body, "fig", 4, fig)
    # tables: prefix caption with its number, drop size/adjustbox wrappers
    def tabcap(m):
        cap, label = m.group(1), m.group(2)
        return f"\\caption{{Table {lab.get(label, '?')}: {cap}}}\\label{{{label}}}"
    body = re.sub(r"\\caption\{((?:(?!\\caption\{).)*?)\}\n\\label\{(tab:[^}]*)\}", tabcap, body, flags=re.S)
    body = re.sub(r"\\begin\{adjustbox\}\{[^}]*\}|\\end\{adjustbox\}", "", body)
    body = re.sub(r"\\(scriptsize|footnotesize|small)\b", "", body)
    body = re.sub(r"\{\s*(\[[^\]]*\])\}", r"\1", body)
    body = body.replace("\\setlength{\\tabcolsep}{4pt}", "")
    body = body.replace("table*", "table").replace("figure*", "figure")
    # references resolved to the PDF's numbers
    body = re.sub(r"\\S~?\\ref\{([^}]*)\}", lambda m: "§" + lab.get(m.group(1), "?"), body)
    body = re.sub(r"\\ref\{([^}]*)\}", lambda m: lab.get(m.group(1), "?"), body)
    body = re.sub(r"\\hyperref\[[^\]]*\]\{([^}]*)\}", r"\1", body)
    body = body.replace("\\citealp{", "\\citet{")
    body = re.sub(r"\\label\{[^}]*\}", "", body)
    body = body.replace("\\begin{abstract}", "\\section*{Abstract}").replace("\\end{abstract}", "")
    body = body.replace("\\mathrm{Hit@5}", "\\mathrm{Hit}\\text{@}5")
    # number the section headings exactly as in the PDF
    sec = sub = 0
    app = False
    lines = []
    for line in body.splitlines():
        if line.startswith("\\section{") and "Reader Prompts" in line:
            app, sec = True, 0
        m = re.match(r"\\section\{(.*)\}", line)
        if m:
            sec, sub = sec + 1, 0
            num = chr(64 + sec) if app else str(sec)
            line = f"\\section*{{{num} {m.group(1)}}}"
        m = re.match(r"\\subsection\{(.*)\}", line)
        if m:
            sub += 1
            line = f"\\subsection*{{{sec}.{sub} {m.group(1)}}}"
        lines.append(line)
    body = "\n".join(lines)

    title = re.search(r"\\title\{(.*?)\}\n", src, re.S).group(1).replace("\\\\", " ").replace("\n", " ")
    doc = ("\\documentclass{article}\n\\usepackage{graphicx}\n"
           f"\\title{{{title}}}\n"
           "\\author{Muhammad Adeel\\\\Zhengzhou University\\\\muhammadadeel167@stu.zzu.edu.cn\\\\adeelcheena791@gmail.com}\n"
           "\\begin{document}\n\\maketitle\n" + body + "\n\\end{document}\n")
    flat = OUT / "paper_flat.tex"
    flat.write_text(doc)
    common = ["pandoc", str(flat), "--citeproc", "--bibliography", str(LATEX / "references.bib"),
              "--metadata", "reference-section-title=References", "--metadata", "link-citations=true", "--standalone"]
    subprocess.run(common + ["-o", str(OUT / "paper_raw.docx")], check=True)
    subprocess.run(common + ["-t", "plain", "--wrap=auto", "--columns=100", "-o", str(OUT / "rag_embed_transfer.txt")],
                   check=True)
    txt = (OUT / "rag_embed_transfer.txt").read_text()
    txt = re.sub(r"\$\$T\(R\) = \\frac\{.*?\}\{.*?\} \.\$\$",
                 "T(R) = [EM(R) - EM(R_1)] / [Hit@5(R) - Hit@5(R_1)]", txt)
    (OUT / "rag_embed_transfer.txt").write_text(txt)
    style_docx(OUT / "paper_raw.docx", OUT / "rag_embed_transfer.docx")
    print("wrote", OUT / "rag_embed_transfer.docx", "and", OUT / "rag_embed_transfer.txt")
    return 0


def style_docx(src: Path, dst: Path):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    d = Document(str(src))
    for s in d.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)
        # page number, bottom centre
        p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); run._r.append(el)
            else:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)
    for st in [st for st in d.styles if st.type == 1]:
        st.font.name = "Times New Roman"
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Times New Roman")
    for st in d.styles:
        if st.type == 1 and st.name in ("Normal", "Body Text", "First Paragraph", "Compact"):
            st.font.size = Pt(11)
    sizes = {"title": 15, "heading 1": 12, "heading 2": 11, "heading 3": 11, "heading 4": 11}
    for st in d.styles:
        if st.type == 1 and st.name and st.name.lower() in sizes:
            st.font.size = Pt(sizes[st.name.lower()])
            st.font.bold = True
            st.font.italic = st.name.lower() in ("heading 2", "heading 3")
    for p in d.paragraphs:
        if p.style.name in ("Title", "Author", "Subtitle", "Date"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    for r in par.runs:
                        r.font.size = Pt(8)
    d.save(str(dst))


if __name__ == "__main__":
    raise SystemExit(main())
