"""Human audit of NQ retrieval misses (design/addendum_v0.4.md, "Human audit of NQ misses"): builds the annotation workbook.

Sample: the NQ generation queries (first 2,255 of the seeded order) that are Hit@5 misses for e5-base-v2 on the
sub-corpus, in query order; 100 drawn with numpy.random.default_rng(20261017).choice(..., 100, replace=False), kept in
draw order. Screen 1 shows the question, the reference answers and the five passages (no reader answer, no automatic
answer-string flag, no gold passage). Screen 2 shows the Qwen reader's answer. The author fills in labels A-D; no AI
labels are used. The key file (automatic flag, EM, gold ids) is written separately and should not be opened before
labelling.

    python3 manuscript_ipm/revision_v04/scripts/make_audit_sheet.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from src.analysis.v03_analysis import Data, jl  # noqa: E402

OUT = ROOT / "annotation/v04_miss_audit"
SEED = 20261017


def main():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    D = Data(ROOT)
    q = D.qids("qwen_nq")
    hit = D.retrieval("qwen_nq", "R5")["hit@5"]
    misses = [x for x, h in zip(q, hit) if h == 0]
    draw = np.random.default_rng(SEED).choice(len(misses), 100, replace=False)
    sample = [misses[i] for i in draw]
    joined = {r["qid"]: r for r in jl(ROOT / "paper_inputs/data/processed/nq_joined.jsonl")}
    gens = {r["qid"]: r for r in jl(ROOT / "paper_inputs/results/raw/generations/main__R5_k5_A1__main.jsonl")}
    ans = {r["qid"]: r for r in jl(ROOT / "paper_inputs/results/processed/answers/main__R5_k5_A1__main.jsonl")}
    texts = {r["pid"]: r for r in jl(ROOT / "paper_inputs/v04_s5nq/v04/audit_texts.jsonl")}
    pa = {r["qid"]: r for r in jl(ROOT / "paper_inputs/v04_s5nq/v04/passage_answers.jsonl")
          if r["corpus"] == "main" and r["retriever"] == "R5"}
    top10 = D.nq_top10["R5"]
    for x in sample:
        assert gens[x]["pids"] == top10[x][:5], x        # the reader saw the run's top 5
        assert not (set(gens[x]["pids"]) & set(D.nq_qrels[x]))

    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    for line in [
        "Human audit of NQ retrieval misses (design addendum v0.4). One annotator: the author. Do not use AI to label.",
        "Do Screen1 for all 100 items first. Do not look at Screen2 or at key.csv before Screen1 is complete.",
        "A: Does any of the five passages contain a correct answer to the question? yes / no / unsure. If yes, give the passage number (1-5) in A_passage.",
        "B: Only if A = yes: is that passage about the question's subject, so that it should count as relevant? yes / no.",
        "Screen2, C: is the reader's answer factually correct for the question? yes / no / unsure (use the references and your own knowledge; note sources in the notes column if you checked one).",
        "Screen2, D: only if C = yes and EM = 0 is shown: is the difference an alias or formatting difference? yes / no.",
        f"Sample: e5-base-v2 Hit@5 misses on the sub-corpus among the 2,255 NQ generation queries ({len(misses)} misses); 100 drawn with numpy default_rng({SEED}).",
    ]:
        ins.append([line])
    ins.column_dimensions["A"].width = 160
    wrap = Alignment(wrap_text=True, vertical="top")
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")

    s1 = wb.create_sheet("Screen1")
    h1 = ["item", "question", "reference answers"] + [f"passage {i}" for i in range(1, 6)] + ["A (yes/no/unsure)", "A_passage (1-5)", "B (yes/no)", "notes"]
    s1.append(h1)
    for i, x in enumerate(sample, 1):
        ps = [f"[{texts[p]['title']}] {texts[p]['text']}" for p in gens[x]["pids"]]
        s1.append([i, joined[x]["question"], " | ".join(joined[x]["answers"]), *ps, "", "", "", ""])
    s2 = wb.create_sheet("Screen2")
    s2.append(["item", "question", "reference answers", "reader answer (Qwen2.5-7B, e5-base-v2 top 5)", "EM", "C (yes/no/unsure)", "D (yes/no)", "notes"])
    for i, x in enumerate(sample, 1):
        s2.append([i, joined[x]["question"], " | ".join(joined[x]["answers"]), gens[x]["output"], int(ans[x]["em"]), "", "", ""])
    for ws, widths, lab in ((s1, [6, 40, 25] + [60] * 5 + [12, 12, 10, 30], [9, 10, 11]), (s2, [6, 45, 30, 40, 6, 12, 10, 30], [6, 7])):
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = w
        for c in ws[1]:
            c.font = head
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
            for j in lab:
                row[j - 1].fill = fill
        ws.freeze_panes = "B2"
    for ws, col in ((s1, "I"), (s2, "F")):
        dv = DataValidation(type="list", formula1='"yes,no,unsure"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{col}2:{col}101")
    for ws, col in ((s1, "K"), (s2, "G")):
        dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{col}2:{col}101")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    s1.add_data_validation(dv); dv.add("J2:J101")
    OUT.mkdir(parents=True, exist_ok=True)
    wb.save(OUT / "miss_audit_v04.xlsx")
    with open(OUT / "key.csv", "w") as f:       # do not open before labelling
        f.write("item,qid,auto_answer_string_in_top5,first_answer_rank,em,gold_pids\n")
        for i, x in enumerate(sample, 1):
            f.write(f"{i},{x},{pa[x]['ans_in_topk']},{pa[x]['first_ans_rank'] or ''},{int(ans[x]['em'])},{' '.join(sorted(D.nq_qrels[x]))}\n")
    (OUT / "sample.json").write_text(json.dumps({"seed": SEED, "n_misses": len(misses), "qids": sample}, indent=1))
    print(len(misses), "misses; sample of", len(sample), "->", OUT)


if __name__ == "__main__":
    main()
