"""Build the IP&M submission package in manuscript_ipm/submission_ipm/.

    python3 manuscript_ipm/revision_v04/scripts/make_submission_package.py

Needs: latexmk (TeX Live), python-docx, LibreOffice (docx -> pdf). Copies the anonymised manuscript PDF and a
compile-tested LaTeX source zip, and writes the title page, highlights, competing-interest declaration and cover
letter as Word files (plus PDFs). Texts of the declarations are the ones in latex/title_page.tex.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[3]
TEX = ROOT / "manuscript_ipm/revision_v04/latex"
OUT = ROOT / "manuscript_ipm/submission_ipm"

TITLE = ("How much retrieval gain reaches the answer? Transfer ratios and retrieval-stage cost of open embedding "
         "models in retrieval-augmented question answering")
AUTHOR = "Muhammad Adeel"
ORCID = "0009-0006-8541-4422"
EMAIL = "muhammadadeel167@stu.zzu.edu.cn"
AFFIL = "Zhengzhou University, Zhengzhou, Henan, China"
REPO = "https://github.com/adeelcheena791/rag-embed-transfer"
COI = ("The author declares that he has no known competing financial interests or personal relationships that could "
       "have appeared to influence the work reported in this paper.")
FUND = ("This research did not receive any specific grant from funding agencies in the public, commercial, or "
        "not-for-profit sectors.")
CREDIT = ("Muhammad Adeel: Conceptualization, Methodology, Software, Investigation, Formal analysis, Data curation, "
          "Visualization, Writing – original draft, Writing – review & editing.")


def doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(12)
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    return d


def to_pdf(p: Path):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(p.parent), str(p)],
                   check=True, capture_output=True)


def expand_ev(s: str) -> str:
    """Replace \\ev{ID} with its value from the numbers files (same values the PDF prints)."""
    vals = {}
    for f in sorted(TEX.glob("numbers*.tex")):
        for m in re.finditer(r"\\csname ev@([^\\]+)\\endcsname\{([^}]*)\}", f.read_text()):
            vals[m.group(1)] = m.group(2)
    s = re.sub(r"\\ev\{([^}]+)\}", lambda m: vals[m.group(1)], s)
    return s.replace("--", "–").replace("\\%", "%")


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)

    # 1. anonymised manuscript PDF
    subprocess.run(["latexmk", "-pdf", "-interaction=nonstopmode", "manuscript.tex"], cwd=TEX, check=True,
                   capture_output=True)
    shutil.copy(TEX / "manuscript.pdf", OUT / "03_Manuscript_anonymised.pdf")

    # 2. LaTeX source, compile-tested in a clean folder
    need = set()
    for line in (TEX / "manuscript.fls").read_text().splitlines():
        if line.startswith("INPUT "):
            f = line[6:].removeprefix("./")
            if not f.startswith("/") and not f.endswith((".aux", ".out", ".abs", ".toc")):
                need.add(f)
    need |= {"references.bib", "cas-model2-names.bst"}
    build = OUT / "_build_test/manuscript_source"
    for f in sorted(need):
        (build / f).parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(TEX / f, build / f)
    with zipfile.ZipFile(OUT / "04_Manuscript_LaTeX_source.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(need):
            z.write(build / f, f)
    subprocess.run(["latexmk", "-pdf", "-interaction=nonstopmode", "manuscript.tex"], cwd=build, check=True,
                   capture_output=True)
    assert (build / "manuscript.pdf").exists()
    shutil.rmtree(OUT / "_build_test")

    # 3. title page (Word + PDF from LaTeX)
    d = doc()
    d.add_heading(TITLE, level=1)
    d.add_paragraph(f"{AUTHOR}\nORCID: {ORCID}\n{AFFIL}\nCorresponding author: {AUTHOR}, {EMAIL}")
    for h, t in (("CRediT authorship contribution statement", CREDIT), ("Declaration of competing interest", COI),
                 ("Funding", FUND),
                 ("Declaration of generative AI and AI-assisted technologies in the manuscript preparation process",
                  "Given in the anonymised manuscript, before the references."),
                 ("Data availability", f"Code, prompts, design documents, raw model outputs and analysis scripts: "
                                       f"{REPO}. An archived copy with a DOI will be deposited on acceptance.")):
        d.add_heading(h, level=2)
        d.add_paragraph(t)
    d.save(OUT / "02_Title_page.docx")
    subprocess.run(["latexmk", "-pdf", "-interaction=nonstopmode", "title_page.tex"], cwd=TEX, check=True,
                   capture_output=True)
    shutil.copy(TEX / "title_page.pdf", OUT / "02_Title_page.pdf")

    # 4. highlights
    ab = (TEX / "sections/00_abstract.tex").read_text()
    hl = re.findall(r"\\item (.*)", re.search(r"\\begin\{highlights\}(.*?)\\end\{highlights\}", ab, re.S).group(1))
    d = doc()
    d.add_heading("Highlights", level=1)
    for h in hl:
        h = expand_ev(h)
        assert len(h) <= 85, h
        d.add_paragraph(h, style="List Bullet")
    d.save(OUT / "05_Highlights.docx")

    # 5. competing interest declaration
    d = doc()
    d.add_heading("Declaration of competing interest", level=1)
    d.add_paragraph(f"Manuscript: {TITLE}")
    d.add_paragraph(COI)
    d.add_paragraph(FUND)
    d.add_paragraph(f"{AUTHOR}, {AFFIL}")
    d.save(OUT / "06_Declaration_of_competing_interest.docx")

    # 6. cover letter
    d = doc()
    d.add_paragraph(f"{AUTHOR}\n{AFFIL}\n{EMAIL}\nORCID: {ORCID}")
    d.add_paragraph("24 October 2026")
    d.add_paragraph("To the Editors-in-Chief,\nInformation Processing & Management")
    d.add_paragraph("Dear Editors,")
    for t in (
        f"I submit the research article “{TITLE}” for consideration in Information Processing & Management.",
        "Dense retrievers for retrieval-augmented generation are usually chosen by their retrieval scores. The "
        "manuscript asks how much of a paired retrieval gain reaches the answer. On Natural Questions and HotpotQA, "
        "it compares BM25 and six open embedding models under two 7–8B readers with corpus, prompt and decoding "
        "held fixed, and reports a paired transfer ratio, the exact-match gain divided by the Hit@5 gain, with "
        "bootstrap intervals. On NQ, dense encoders passed " + expand_ev(r"\ev{T.lnq.R4}--\ev{T.lnq.R3}") + " exact-match points per Hit@5 point, and a "
        "registered replication on the full 2.68-million-passage corpus agreed for the encoders it covered. On "
        "HotpotQA, the ratios were lower once retrieval was scored by complete evidence. The paper also "
        "reports the measured retrieval-stage cost of each encoder on a free GPU.",
        "The work fits the journal’s interest in the evaluation of information retrieval and question-answering "
        "systems. Its predictions were registered in dated design documents before the second and third study "
        "phases, all deviations are listed, and every reported number is generated by script from the committed "
        "outputs. Code, prompts, raw outputs and analysis scripts are available in the repository linked on the "
        "title page.",
        "The manuscript is original, has not been published, and is not under consideration elsewhere. It has not "
        "been posted as a preprint. I am the sole author and have no competing interests; the work received no "
        "specific funding. The use of generative AI tools is declared in the manuscript. In line with the journal’s "
        "double-anonymised review, the manuscript file contains no author information; author details are in the "
        "separate title page.",
        "Thank you for considering this submission.",
    ):
        d.add_paragraph(t)
    d.add_paragraph(f"Sincerely,\n{AUTHOR}")
    d.save(OUT / "01_Cover_letter.docx")

    for f in ("01_Cover_letter.docx", "05_Highlights.docx", "06_Declaration_of_competing_interest.docx"):
        to_pdf(OUT / f)
    print("\n".join(sorted(p.name for p in OUT.iterdir())))


if __name__ == "__main__":
    main()
